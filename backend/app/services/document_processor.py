import io
from typing import Optional, List
import PyPDF2
from docx import Document as DocxDocument
import openpyxl
from app.core.config import settings

class DocumentProcessor:
    """Extract text from various document formats"""

    @staticmethod
    def extract_from_pdf(file_bytes: bytes) -> str:
        try:
            pdf = PyPDF2.PdfReader(io.BytesIO(file_bytes))
            text = ""
            for page in pdf.pages:
                text += page.extract_text() or ""
            return text
        except Exception as e:
            return f"Error reading PDF: {str(e)}"

    @staticmethod
    def extract_from_docx(file_bytes: bytes) -> str:
        try:
            doc = DocxDocument(io.BytesIO(file_bytes))
            text = "
".join([paragraph.text for paragraph in doc.paragraphs])
            return text
        except Exception as e:
            return f"Error reading DOCX: {str(e)}"

    @staticmethod
    def extract_from_txt(file_bytes: bytes) -> str:
        try:
            return file_bytes.decode("utf-8")
        except:
            try:
                return file_bytes.decode("latin-1")
            except Exception as e:
                return f"Error reading text: {str(e)}"

    @staticmethod
    def extract_from_xlsx(file_bytes: bytes) -> str:
        try:
            wb = openpyxl.load_workbook(io.BytesIO(file_bytes))
            text = ""
            for sheet in wb.worksheets:
                text += f"Sheet: {sheet.title}
"
                for row in sheet.iter_rows(values_only=True):
                    text += " | ".join([str(cell) for cell in row if cell is not None]) + "
"
            return text
        except Exception as e:
            return f"Error reading XLSX: {str(e)}"

    @staticmethod
    def extract(file_bytes: bytes, mime_type: str) -> str:
        if mime_type == "application/pdf":
            return DocumentProcessor.extract_from_pdf(file_bytes)
        elif mime_type in ["application/vnd.openxmlformats-officedocument.wordprocessingml.document", "application/msword"]:
            return DocumentProcessor.extract_from_docx(file_bytes)
        elif mime_type in ["text/plain", "text/markdown", "text/csv"]:
            return DocumentProcessor.extract_from_txt(file_bytes)
        elif mime_type in ["application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", "application/vnd.ms-excel"]:
            return DocumentProcessor.extract_from_xlsx(file_bytes)
        else:
            return f"Unsupported file type: {mime_type}"

    @staticmethod
    def chunk_text(text: str, chunk_size: int = 1000, overlap: int = 200) -> List[str]:
        """Split text into overlapping chunks for RAG"""
        chunks = []
        start = 0
        while start < len(text):
            end = min(start + chunk_size, len(text))
            # Try to end at a sentence boundary
            if end < len(text):
                for sep in [".
", ". ", "!
", "! ", "?
", "? ", "

", "
"]:
                    pos = text.rfind(sep, start, end)
                    if pos > start + chunk_size // 2:
                        end = pos + len(sep)
                        break
            chunks.append(text[start:end].strip())
            start = end - overlap
            if start < 0:
                start = 0
        return chunks

doc_processor = DocumentProcessor()
